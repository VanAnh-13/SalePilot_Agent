"""Deterministic import of Spec_cate_gia.xlsx into snapshot (+ optional Mongo).

Produces a canonical sorted JSON snapshot and writes hashes into
``experiments/manifest.json`` when present.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import tempfile
from pathlib import Path
from typing import Any

import openpyxl

from app.catalog import repository
from app.catalog.categories import BY_SHEET, EXPECTED_SHEETS, Category, normalize_workbook_product
from app.config import get_settings

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_EXCEL = ROOT / "Spec_cate_gia.xlsx"
MANIFEST_PATH = ROOT / "experiments" / "manifest.json"
RESEARCH_SNAPSHOT = ROOT / "backend" / "data" / "research" / "catalog_workbook_snapshot.json"
RESEARCH_MONGO_COLLECTION = "research_catalog_workbook_products"
COMMON_HEADERS = {"sku"}


def _sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _canonical_json_bytes(products: list[dict[str, Any]]) -> bytes:
    return (json.dumps(products, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False) + "\n").encode(
        "utf-8"
    )


def _rows(ws, sheet_name: str) -> list[dict[str, Any]]:
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        raise ValueError(f"Sheet {sheet_name!r} has no header row")
    header = [str(c).strip() if c is not None else "" for c in header_row]
    if any(not h for h in header):
        raise ValueError(f"Sheet {sheet_name!r} has blank header cells")
    if len(header) != len(set(header)):
        raise ValueError(f"Sheet {sheet_name!r} has duplicate headers")
    lower = {h.casefold() for h in header}
    if "sku" not in lower:
        raise ValueError(f"Sheet {sheet_name!r} missing required column 'sku'")

    rows: list[dict[str, Any]] = []
    for offset, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if values is None or all(v in (None, "") for v in values):
            continue
        row = {header[i]: values[i] for i in range(min(len(header), len(values)))}
        row["__row__"] = offset
        rows.append(row)
    return rows


def build_catalog(excel: Path, *, require_all_sheets: bool = True) -> tuple[list[dict[str, Any]], dict[str, dict[str, int]]]:
    wb = openpyxl.load_workbook(excel, read_only=True, data_only=True)
    try:
        sheet_names = list(wb.sheetnames)
        present = {name for name in sheet_names if name in BY_SHEET}
        missing = [name for name in EXPECTED_SHEETS if name not in present]
        if require_all_sheets and missing:
            raise ValueError(f"Missing expected sheets: {missing}")

        products: list[dict[str, Any]] = []
        stats: dict[str, dict[str, int]] = {}
        seen_skus: dict[str, tuple[str, int]] = {}

        # Process in registry order for stable intermediate behavior.
        for sheet in EXPECTED_SHEETS:
            if sheet not in present:
                continue
            cat: Category = BY_SHEET[sheet]
            source = f"spec_cate_gia.xlsx:{sheet}"
            total = priced = 0
            for row in _rows(wb[sheet], sheet):
                sku = str(row.get("sku") or row.get("SKU") or "").strip()
                if not sku:
                    raise ValueError(f"Blank sku at {sheet} row {row['__row__']}")
                if sku in seen_skus:
                    prev_sheet, prev_row = seen_skus[sku]
                    raise ValueError(
                        f"Duplicate sku {sku!r}: {prev_sheet} row {prev_row} and {sheet} row {row['__row__']}"
                    )
                doc = normalize_workbook_product(cat, row, int(row["__row__"]), source)
                seen_skus[sku] = (sheet, int(row["__row__"]))
                products.append(doc)
                total += 1
                priced += int(bool(doc.get("has_current_price")))
            stats[cat.slug] = {"code": cat.code, "total": total, "priced": priced}
    finally:
        wb.close()

    products.sort(key=lambda d: (int(d.get("category_code") or 0), str(d.get("sku") or "")))
    return products, stats


def save_canonical_snapshot(products: list[dict[str, Any]], path: Path | None = None) -> tuple[Path, str]:
    payload = _canonical_json_bytes(products)
    digest = _sha256_bytes(payload)
    target = path or RESEARCH_SNAPSHOT
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("wb", delete=False, dir=str(target.parent)) as tmp:
        tmp.write(payload)
        tmp_path = Path(tmp.name)
    tmp_path.replace(target)
    return target, digest


def write_mongo(products: list[dict[str, Any]]) -> int:
    from pymongo import ASCENDING, UpdateOne

    settings = get_settings()
    client = repository.mongo_client(timeout_ms=4000)
    coll = client[settings.mongodb_db][RESEARCH_MONGO_COLLECTION]
    coll.create_index([("sku", ASCENDING)], unique=True)
    coll.create_index([("category_code", ASCENDING)])
    coll.create_index([("brand", ASCENDING)])
    coll.create_index([("price_vnd", ASCENDING)])
    ops = [UpdateOne({"sku": doc["sku"]}, {"$set": doc}, upsert=True) for doc in products]
    for start in range(0, len(ops), 1000):
        coll.bulk_write(ops[start : start + 1000], ordered=False)
    current = {doc["sku"] for doc in products}
    coll.delete_many({"sku": {"$nin": list(current)}})
    count = coll.count_documents({})
    client.close()
    return count


def validate_import_rights(
    excel: Path,
    snapshot_path: Path,
    *,
    confirm_publication_rights: bool = False,
    allow_partial_sheets: bool = False,
) -> None:
    """Validate rights and declared source path before any write occurs."""
    if allow_partial_sheets:
        forbidden = {repository._snapshot_path().resolve(), RESEARCH_SNAPSHOT.resolve()}  # noqa: SLF001
        if snapshot_path.resolve() in forbidden:
            raise ValueError("partial import snapshot must not overwrite runtime or canonical research data")
        return
    if not MANIFEST_PATH.is_file():
        raise ValueError("research manifest is required before a full workbook import")
    data = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    catalog = data.get("catalog") or {}
    denied = {"denied", "revoked", "not_allowed"}
    if catalog.get("rights_status") in denied:
        raise ValueError("manifest catalog rights are denied")
    if not confirm_publication_rights and catalog.get("rights_status") != "owner_confirmed_publication":
        raise ValueError(
            "manifest catalog rights are not publication-approved; "
            "pass --confirm-publication-rights after owner approval"
        )

    declared = catalog.get("source_path")
    declared_path = (ROOT / declared).resolve() if isinstance(declared, str) else None
    if declared_path is None or declared_path != excel.resolve():
        raise ValueError(
            "import source does not match manifest catalog.source_path: "
            f"{excel} != {declared}"
        )
    declared_snapshot = catalog.get("snapshot_path")
    manifest_snapshot = (
        (ROOT / declared_snapshot).resolve() if isinstance(declared_snapshot, str) else None
    )
    if manifest_snapshot is None or manifest_snapshot != snapshot_path.resolve():
        raise ValueError("import snapshot target does not match manifest catalog.snapshot_path")
    for asset in data.get("assets") or []:
        if asset.get("rights_status") in denied:
            raise ValueError(f"asset rights are denied: {asset.get('id')}")
        if asset.get("id") == "catalog_workbook":
            if (
                not confirm_publication_rights
                and asset.get("rights_status") != "owner_confirmed_publication"
            ):
                raise ValueError("catalog_workbook asset rights are not publication-approved")
            asset_path = asset.get("path")
            resolved_asset = (ROOT / asset_path).resolve() if isinstance(asset_path, str) else None
            if resolved_asset != excel.resolve():
                raise ValueError("catalog_workbook asset path does not match import source")


def update_manifest(
    *,
    raw_hash: str,
    normalized_hash: str,
    total: int,
    stats: dict[str, dict[str, int]],
    source_path: str,
    confirm_publication_rights: bool = False,
) -> None:
    if not MANIFEST_PATH.is_file():
        return
    data = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    catalog = data.setdefault("catalog", {})

    # Never turn a source with denied or unconfirmed rights into a publication
    # asset as a side effect of importing it. The explicit flag is only for a
    # pending declaration; denied rights always fail closed.
    if catalog.get("rights_status") in {"denied", "revoked", "not_allowed"}:
        raise ValueError("manifest catalog rights are denied; import cannot mark it ready")
    if confirm_publication_rights:
        catalog["rights_status"] = "owner_confirmed_publication"
        for asset in data.get("assets") or []:
            if asset.get("id") == "catalog_workbook":
                if asset.get("rights_status") in {"denied", "revoked", "not_allowed"}:
                    raise ValueError("manifest workbook rights are denied")
                asset["rights_status"] = "owner_confirmed_publication"
    elif catalog.get("rights_status") != "owner_confirmed_publication":
        raise ValueError(
            "manifest catalog rights are not publication-approved; "
            "pass --confirm-publication-rights after owner approval"
        )

    catalog["source_path"] = source_path
    catalog["source_present"] = True
    catalog["raw_source_hash"] = raw_hash
    catalog["normalized_catalog_hash"] = normalized_hash
    catalog["expected_total_products"] = total
    catalog["expected_categories"] = len(stats)
    catalog["per_category"] = stats
    data["status"] = "ready"
    data["blockers"] = []
    for asset in data.get("assets") or []:
        if asset.get("id") == "catalog_workbook":
            asset["hash"] = raw_hash
            asset["rights_status"] = "owner_confirmed_publication"
        if asset.get("id") == "catalog_snapshot":
            asset["hash"] = normalized_hash
            asset["rights_status"] = "derived_ready"
    MANIFEST_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", type=Path, default=DEFAULT_EXCEL)
    parser.add_argument("--snapshot-only", action="store_true")
    parser.add_argument("--allow-partial-sheets", action="store_true", help="For synthetic fixtures/tests")
    parser.add_argument("--no-manifest-update", action="store_true")
    parser.add_argument(
        "--snapshot-out",
        type=Path,
        default=None,
        help="Output path; required for partial engineering imports",
    )
    parser.add_argument(
        "--confirm-publication-rights",
        action="store_true",
        help="Explicitly promote a pending owner-rights declaration before marking the manifest ready",
    )
    args = parser.parse_args(argv)

    if args.allow_partial_sheets and not args.snapshot_only:
        raise SystemExit("--allow-partial-sheets requires --snapshot-only; refusing Mongo writes/pruning")
    if args.allow_partial_sheets and not args.no_manifest_update:
        raise SystemExit("--allow-partial-sheets requires --no-manifest-update; refusing manifest promotion")
    if args.allow_partial_sheets and args.snapshot_out is None:
        raise SystemExit("--allow-partial-sheets requires an explicit --snapshot-out")

    excel = args.excel if args.excel.is_absolute() else (Path.cwd() / args.excel).resolve()
    if not excel.exists():
        raise SystemExit(f"Excel not found: {excel}")
    snapshot_target = args.snapshot_out or RESEARCH_SNAPSHOT
    if not snapshot_target.is_absolute():
        snapshot_target = (Path.cwd() / snapshot_target).resolve()

    try:
        validate_import_rights(
            excel,
            snapshot_target,
            confirm_publication_rights=args.confirm_publication_rights,
            allow_partial_sheets=args.allow_partial_sheets,
        )
    except ValueError as exc:
        raise SystemExit(str(exc)) from exc

    raw_hash = _sha256_file(excel)
    products, stats = build_catalog(excel, require_all_sheets=not args.allow_partial_sheets)
    if not products:
        raise SystemExit("No products parsed")

    snapshot, norm_hash = save_canonical_snapshot(products, path=snapshot_target)
    print(f"Snapshot: {snapshot} products={len(products)} sha256={norm_hash}")
    print(f"Raw source sha256={raw_hash}")

    if not args.snapshot_only:
        try:
            count = write_mongo(products)
            print(f"MongoDB collection={RESEARCH_MONGO_COLLECTION} docs={count}")
        except Exception as exc:  # pragma: no cover
            print(f"MongoDB write failed: {exc}")
            if not args.allow_partial_sheets:
                # Fail closed for real imports unless snapshot-only.
                raise SystemExit(2) from exc

    if not args.no_manifest_update:
        rel = str(excel.relative_to(ROOT)) if excel.is_relative_to(ROOT) else str(excel)
        try:
            update_manifest(
                raw_hash=raw_hash,
                normalized_hash=norm_hash,
                total=len(products),
                stats=stats,
                source_path=rel.replace("\\", "/"),
                confirm_publication_rights=args.confirm_publication_rights,
            )
        except ValueError as exc:
            raise SystemExit(str(exc)) from exc
        print(f"Updated manifest: {MANIFEST_PATH}")

    for slug, s in sorted(stats.items(), key=lambda kv: -kv[1]["total"]):
        print(f"  {slug:18s} code={s['code']:<4} total={s['total']:<5} priced={s['priced']}")
    print(f"TOTAL {len(products)} products / {len(stats)} categories")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
