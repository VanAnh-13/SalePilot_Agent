"""Synthetic workbook import contract tests (no real Spec_cate_gia.xlsx required)."""

from __future__ import annotations

import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from app.catalog.categories import (
    BY_SHEET,
    CATEGORIES,
    EXPECTED_SHEETS,
    Spec,
    _apply_spec,
    detect_category,
    detect_unsupported,
)
from scripts.import_spec_catalog import build_catalog, main, save_canonical_snapshot
import scripts.import_spec_catalog as importer


def _write_min_workbook(path: Path, *, sheets: list[str] | None = None, extra_row: dict | None = None) -> None:
    wb = Workbook()
    # remove default sheet
    wb.remove(wb.active)
    use_sheets = sheets or list(EXPECTED_SHEETS)
    for idx, sheet in enumerate(use_sheets):
        ws = wb.create_sheet(sheet)
        headers = ["sku", "brand", "giá gốc", "giá khuyến mãi", "Số người sử dụng", "Dung tích sử dụng", "Ngang", "Công nghệ tiết kiệm điện", "Lấy nước ngoài", "Phạm vi làm lạnh hiệu quả", "Khối lượng giặt", "Nghe gọi", "SIM", "RAM", "Kích thước"]
        ws.append(headers)
        sku = f"SKU{idx+1:04d}"
        row = {
            "sku": sku,
            "brand": "TestBrand",
            "giá gốc": 12_000_000,
            "giá khuyến mãi": 10_000_000,
            "Số người sử dụng": "3 - 4 người",
            "Dung tích sử dụng": "300 lít",
            "Ngang": "60",
            "Công nghệ tiết kiệm điện": "Inverter",
            "Lấy nước ngoài": "Có",
            "Phạm vi làm lạnh hiệu quả": "Từ 15 - 20m²",
            "Khối lượng giặt": "9 kg",
            "Nghe gọi": "Có",
            "SIM": "Có",
            "RAM": "16 GB",
            "Kích thước": "27 inch",
        }
        if extra_row and sheet == use_sheets[0]:
            row.update(extra_row)
        ws.append([row.get(h) for h in headers])
    wb.save(path)


class ImportSpecCatalogTests(unittest.TestCase):
    def test_registry_has_14_unique_sheets(self):
        self.assertEqual(len(EXPECTED_SHEETS), 14)
        self.assertEqual(len(BY_SHEET), 14)
        self.assertEqual(len(set(EXPECTED_SHEETS)), 14)
        self.assertTrue(all(slot.question for cat in CATEGORIES for slot in cat.slots if slot.primary))

    def test_present_mode_preserves_explicit_negative(self):
        spec = Spec("feature", "Feature", "present")
        self.assertFalse(_apply_spec(spec, "Không"))
        self.assertFalse(_apply_spec(spec, "Khong co"))
        self.assertTrue(_apply_spec(spec, "Theo dõi sức khỏe"))

    def test_build_catalog_deterministic_and_sorted(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            xlsx = root / "mini.xlsx"
            _write_min_workbook(xlsx)
            products_a, stats_a = build_catalog(xlsx, require_all_sheets=True)
            products_b, stats_b = build_catalog(xlsx, require_all_sheets=True)
            self.assertEqual(len(products_a), 14)
            self.assertEqual(len(stats_a), 14)
            self.assertEqual(products_a, products_b)
            self.assertEqual(stats_a, stats_b)
            skus = [p["sku"] for p in products_a]
            self.assertEqual(skus, sorted(skus, key=lambda s: (next(p["category_code"] for p in products_a if p["sku"] == s), s)))
            codes = [p["category_code"] for p in products_a]
            self.assertEqual(codes, sorted(codes))
            path, digest1 = save_canonical_snapshot(products_a, path=root / "snap.json")
            raw = path.read_bytes()
            self.assertEqual(hashlib.sha256(raw).hexdigest(), digest1)
            _, digest2 = save_canonical_snapshot(products_b, path=root / "snap2.json")
            self.assertEqual(digest1, digest2)
            # mutation changes hash
            mutated = json.loads(path.read_text(encoding="utf-8"))
            mutated[0]["price_vnd"] = 1
            mut_path = root / "mut.json"
            mut_path.write_bytes(
                (json.dumps(mutated, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n").encode("utf-8")
            )
            self.assertNotEqual(hashlib.sha256(mut_path.read_bytes()).hexdigest(), digest1)

    def test_duplicate_sku_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            xlsx = Path(tmp) / "dup.xlsx"
            wb = Workbook()
            wb.remove(wb.active)
            for sheet in EXPECTED_SHEETS[:2]:
                ws = wb.create_sheet(sheet)
                ws.append(["sku", "brand", "giá gốc"])
                ws.append(["SAME", "A", 1000000])
            wb.save(xlsx)
            with self.assertRaises(ValueError):
                build_catalog(xlsx, require_all_sheets=False)

    def test_missing_sheet_fails_when_required(self):
        with tempfile.TemporaryDirectory() as tmp:
            xlsx = Path(tmp) / "partial.xlsx"
            _write_min_workbook(xlsx, sheets=list(EXPECTED_SHEETS[:3]))
            with self.assertRaises(ValueError):
                build_catalog(xlsx, require_all_sheets=True)

    def test_partial_cli_requires_snapshot_only_without_manifest_update(self):
        with self.assertRaisesRegex(SystemExit, "requires --snapshot-only"):
            main(["--allow-partial-sheets"])
        with self.assertRaisesRegex(SystemExit, "requires --no-manifest-update"):
            main(["--allow-partial-sheets", "--snapshot-only"])
        with self.assertRaisesRegex(SystemExit, "requires an explicit --snapshot-out"):
            main(
                [
                    "--allow-partial-sheets",
                    "--snapshot-only",
                    "--no-manifest-update",
                ]
            )

    def test_denied_rights_fail_before_snapshot_or_database_write(self):
        with tempfile.TemporaryDirectory(dir=importer.ROOT) as tmp:
            root = Path(tmp)
            excel = root / "denied.xlsx"
            excel.write_bytes(b"not parsed because preflight must fail first")
            rel = excel.relative_to(importer.ROOT).as_posix()
            manifest = root / "manifest.json"
            manifest.write_text(
                json.dumps(
                    {
                        "catalog": {
                            "source_path": rel,
                            "rights_status": "denied",
                        },
                        "assets": [
                            {
                                "id": "catalog_workbook",
                                "path": rel,
                                "rights_status": "denied",
                            }
                        ],
                    }
                ),
                encoding="utf-8",
            )
            with (
                patch.object(importer, "MANIFEST_PATH", manifest),
                patch.object(importer, "save_canonical_snapshot") as save_snapshot,
                patch.object(importer, "write_mongo") as write_mongo,
            ):
                with self.assertRaisesRegex(SystemExit, "rights are denied"):
                    main(["--excel", str(excel)])
                save_snapshot.assert_not_called()
                write_mongo.assert_not_called()

    def test_research_storage_isolated_and_manifest_required(self):
        self.assertNotEqual(
            importer.RESEARCH_SNAPSHOT.resolve(),
            importer.repository._snapshot_path().resolve(),  # noqa: SLF001
        )
        self.assertNotEqual(
            importer.RESEARCH_MONGO_COLLECTION,
            importer.get_settings().mongodb_products_collection,
        )
        with tempfile.TemporaryDirectory(dir=importer.ROOT) as tmp:
            root = Path(tmp)
            excel = root / "source.xlsx"
            excel.write_bytes(b"source")
            missing_manifest = root / "missing-manifest.json"
            with patch.object(importer, "MANIFEST_PATH", missing_manifest):
                with self.assertRaisesRegex(ValueError, "manifest is required"):
                    importer.validate_import_rights(
                        excel,
                        importer.RESEARCH_SNAPSHOT,
                    )

    def test_laptop_unsupported_and_fridge_detected(self):
        self.assertIsNone(detect_category("tư vấn laptop 20 triệu"))
        unsupported = detect_unsupported("tư vấn laptop 20 triệu")
        self.assertIsNotNone(unsupported)
        self.assertEqual(unsupported[0], "laptop")
        cat = detect_category("tủ lạnh gia đình 4 người")
        self.assertIsNotNone(cat)
        self.assertEqual(cat.slug, "tu_lanh")
        self.assertEqual(cat.code, 38)


if __name__ == "__main__":
    raise SystemExit(unittest.main())
